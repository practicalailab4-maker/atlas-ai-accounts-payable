import sqlite3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment

from backend.database import DATA_DIR, DATABASE_PATH


EXPORT_DIR = DATA_DIR / "exports"
EXCEL_PATH = EXPORT_DIR / "atlas_invoice_register.xlsx"


def create_excel_workbook():
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()

    register_sheet = workbook.active
    register_sheet.title = "Invoice Register"

    register_sheet.append([
        "Database ID",
        "Vendor Name",
        "Invoice Number",
        "Invoice Date",
        "Due Date",
        "Bill To",
        "Currency",
        "Subtotal",
        "Tax",
        "Total",
        "Status",
        "Source File",
        "Created At",
    ])

    with sqlite3.connect(DATABASE_PATH) as connection:
        invoices = connection.execute(
            """
            SELECT
                id, vendor_name, invoice_number, invoice_date,
                due_date, bill_to, currency, subtotal, tax,
                total, status, source_file, created_at
            FROM invoices
            ORDER BY id
            """
        ).fetchall()

    for invoice in invoices:
        invoice_row = list(invoice)

        for column_index in (7, 8, 9):
            invoice_row[column_index] = float(invoice_row[column_index])

        register_sheet.append(invoice_row)

    for row in register_sheet.iter_rows(min_row=2):
        for column_index in (7, 8, 9):
            row[column_index].number_format = '"$"#,##0.00'

    line_items_sheet = workbook.create_sheet("Line Items")

    line_items_sheet.append([
        "Line Item ID",
        "Invoice ID",
        "Description",
        "Quantity",
        "Unit Price",
        "Amount",
    ])

    with sqlite3.connect(DATABASE_PATH) as connection:
        line_items = connection.execute(
            """
            SELECT
                id, invoice_id, description,
                quantity, unit_price, amount
            FROM invoice_line_items
            ORDER BY invoice_id, id
            """
        ).fetchall()

    for item in line_items:
        item_row = list(item)

        for column_index in (3, 4, 5):
            item_row[column_index] = float(item_row[column_index])

        line_items_sheet.append(item_row)

    for row in line_items_sheet.iter_rows(min_row=2):
        for column_index in (4, 5):
            row[column_index].number_format = '"$"#,##0.00'

    exceptions_sheet = workbook.create_sheet("Exceptions")

    exceptions_sheet.append([
        "Exception ID",
        "Invoice ID",
        "Issue Message",
        "Created At",
    ])

    with sqlite3.connect(DATABASE_PATH) as connection:
        exceptions = connection.execute(
            """
            SELECT
                id, invoice_id, issue_message, created_at
            FROM invoice_validation_issues
            ORDER BY invoice_id, id
            """
        ).fetchall()

    for exception in exceptions:
        exceptions_sheet.append(exception)

    summary_sheet = workbook.create_sheet("Summary")

    summary_sheet.append(["Metric", "Value"])
    with sqlite3.connect(DATABASE_PATH) as connection:
        total_invoices = connection.execute(
            "SELECT COUNT(*) FROM invoices"
        ).fetchone()[0]

    summary_sheet.append(["Total Invoices", total_invoices])
    with sqlite3.connect(DATABASE_PATH) as connection:
        needs_review = connection.execute(
            "SELECT COUNT(*) FROM invoices WHERE status = ?",
            ("needs_review",),
        ).fetchone()[0]

    summary_sheet.append(["Needs Review", needs_review])

    with sqlite3.connect(DATABASE_PATH) as connection:
        approved_count = connection.execute(
            "SELECT COUNT(*) FROM invoices WHERE status = ?",
            ("approved",),
        ).fetchone()[0]

    summary_sheet.append(["Approved", approved_count])
    with sqlite3.connect(DATABASE_PATH) as connection:
        rejected_count = connection.execute(
            "SELECT COUNT(*) FROM invoices WHERE status = ?",
            ("rejected",),
        ).fetchone()[0]

    summary_sheet.append(["Rejected", rejected_count])
    with sqlite3.connect(DATABASE_PATH) as connection:
        validation_issue_count = connection.execute(
            "SELECT COUNT(*) FROM invoice_validation_issues"
        ).fetchone()[0]

    summary_sheet.append(["Validation Issues", validation_issue_count])


    from decimal import Decimal

    with sqlite3.connect(DATABASE_PATH) as connection:
        invoice_totals = connection.execute(
            "SELECT total FROM invoices"
        ).fetchall()

    total_invoice_value = sum(
        (Decimal(row[0]) for row in invoice_totals),
        Decimal("0.00"),
    )

    summary_sheet.append([
        "Total Invoice Value",
        float(total_invoice_value),
    ])

    summary_sheet["B7"].number_format = '"$"#,##0.00'

    
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="172554",
    )

    header_font = Font(
        name="Aptos",
        size=11,
        bold=True,
        color="FFFFFF",
    )

    alternate_row_fill = PatternFill(
        fill_type="solid",
        fgColor="EFF6FF",
    )


    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="left",
                vertical="center",
            )

        sheet.row_dimensions[1].height = 28

        for row in sheet.iter_rows(min_row=2):
            if row[0].row % 2 == 0:
                for cell in row:
                    cell.fill = alternate_row_fill

        for column_cells in sheet.columns:
            column_letter = get_column_letter(column_cells[0].column)

            max_length = max(
                len(str(cell.value)) if cell.value is not None else 0
                for cell in column_cells
            )

            sheet.column_dimensions[column_letter].width = min(
                max(max_length + 3, 14),
                55,
            )

    workbook.save(EXCEL_PATH)
    return EXCEL_PATH